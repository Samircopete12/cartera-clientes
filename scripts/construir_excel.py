#!/usr/bin/env python3
"""Construye el workbook Cartera_20_Clientes.xlsx con fórmulas y dashboard."""

from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path("/home/rimas/cartera-clientes")
DATOS = BASE / "datos"
SALIDA = BASE / "Cartera_20_Clientes.xlsx"

NAVY = "16324F"
TEAL = "2A7F8F"
GOLD = "C9A227"
LIGHT = "EEF3F7"
LIGHT2 = "F7F9FB"
GRAY = "D5DEE7"
WHITE = "FFFFFF"
RED = "C0392B"
GREEN = "1E8449"
AMBER = "B7950B"
ORANGE = "D35400"
SOFT_RED = "FADBD8"
SOFT_GREEN = "D5F5E3"
SOFT_AMBER = "FCF3CF"
SOFT_BLUE = "D6EAF8"
SOFT_ORANGE = "FDEBD0"

MONEY = '"$"#,##0;[Red]-"$"#,##0'
PCT = "0.0%"
DATE = "dd/mm/yyyy"
INT = "#,##0"

THIN = Side(style="thin", color=GRAY)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def leer(nombre: str) -> list[dict]:
    with (DATOS / nombre).open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def fecha(valor: str) -> date:
    return date.fromisoformat(valor)


def titulo(ws, texto: str, rango: str, size: int = 18, color: str = NAVY):
    ws.merge_cells(rango)
    cell = ws[rango.split(":")[0]]
    cell.value = texto
    cell.font = Font(name="Calibri", size=size, bold=True, color=color)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[cell.row].height = max(24, size + 10)


def subtitulo(ws, texto: str, rango: str):
    ws.merge_cells(rango)
    cell = ws[rango.split(":")[0]]
    cell.value = texto
    cell.font = Font(name="Calibri", size=10, italic=True, color="5F6F7F")
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def encabezado(ws, fila: int, columnas: int):
    for col in range(1, columnas + 1):
        c = ws.cell(row=fila, column=col)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[fila].height = 32


def ajustar_anchos(ws, anchos: dict[str, int]):
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho


def formato_rango(ws, rango: str, fmt: str, alineacion: str = "right"):
    for row in ws[rango]:
        for c in row:
            c.number_format = fmt
            c.alignment = Alignment(horizontal=alineacion, vertical="center")


def borde_rango(ws, rango: str):
    for row in ws[rango]:
        for c in row:
            c.border = BORDER


def fila_total(ws, fila: int, columnas: int, etiqueta: str):
    ws.cell(row=fila, column=1, value=etiqueta)
    ws.cell(row=fila, column=1).font = Font(bold=True, color=NAVY)
    for col in range(1, columnas + 1):
        c = ws.cell(row=fila, column=col)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.border = BORDER
        c.font = Font(bold=True, color=NAVY)


def crear_hoja_datos(wb: Workbook, nombre: str, tab: str):
    ws = wb.create_sheet(nombre)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    return ws


def main() -> None:
    clientes = leer("clientes.csv")
    facturas = leer("facturas.csv")
    nc = leer("notas_credito.csv")
    nd = leer("notas_debito.csv")
    pagos = leer("pagos.csv")
    aplicaciones = leer("aplicacion_pagos.csv")
    extracto = leer("extracto_bancario.csv")

    wb = Workbook()
    ws_ini = wb.active
    ws_ini.title = "Inicio"
    ws_ini.sheet_properties.tabColor = GOLD
    ws_ini.sheet_view.showGridLines = False

    ws_dash = crear_hoja_datos(wb, "Dashboard", GOLD)
    ws_cart = crear_hoja_datos(wb, "Cartera", TEAL)
    ws_aging = crear_hoja_datos(wb, "Aging", TEAL)
    ws_conc = crear_hoja_datos(wb, "Conciliacion", TEAL)
    ws_pagos = crear_hoja_datos(wb, "Pagos", "6C8EBF")
    ws_notas = crear_hoja_datos(wb, "Notas", "6C8EBF")
    ws_fact = crear_hoja_datos(wb, "Facturas", "6C8EBF")
    ws_cli = crear_hoja_datos(wb, "Clientes", "6C8EBF")
    ws_apl = crear_hoja_datos(wb, "Aplicacion", "6C8EBF")
    ws_ext = crear_hoja_datos(wb, "Extracto", "6C8EBF")
    ws_val = crear_hoja_datos(wb, "Validacion", "7F8C8D")

    # ---------------------------------------------------------------
    # Inicio
    # ---------------------------------------------------------------
    titulo(ws_ini, "Cartera de clientes — caso de 20 clientes", "A1:I1", 22)
    subtitulo(
        ws_ini,
        "Empresa ficticia: Distribuciones Andinas S.A.S. · Corte: 30/09/2026 · Datos 100% sintéticos",
        "A2:I2",
    )
    ws_ini["A4"] = "Fecha de corte"
    ws_ini["A4"].font = Font(bold=True, color=NAVY)
    ws_ini["B4"] = date(2026, 9, 30)
    ws_ini["B4"].number_format = DATE
    ws_ini["B4"].fill = PatternFill("solid", fgColor=SOFT_AMBER)
    ws_ini["B4"].border = BORDER
    ws_ini["D4"] = "Moneda"
    ws_ini["D4"].font = Font(bold=True, color=NAVY)
    ws_ini["E4"] = "COP"
    ws_ini["E4"].border = BORDER

    textos = [
        ("A6", "Objetivo del caso"),
        (
            "A7",
            "Gestionar la cartera de 20 clientes: emitir y registrar notas crédito y débito, aplicar pagos, "
            "cruzar saldos entre clientes y empresa, consultar el extracto bancario y construir el aging de cartera.",
        ),
        ("A9", "Qué contiene el workbook"),
        ("A10", "Dashboard: KPIs, aging, top 5 clientes y conciliación bancaria."),
        ("A11", "Cartera: saldo por factura después de notas y pagos aplicados."),
        ("A12", "Aging: saldo por cliente y por rango de vencimiento (vigente, 0-30, 31-60, 61-90, >90)."),
        ("A13", "Conciliacion: pagos registrados contra extracto bancario, con diferencias identificadas."),
        ("A14", "Pagos / Aplicacion: recaudos, aplicación a facturas y saldos no aplicados."),
        ("A15", "Notas / Facturas / Clientes / Extracto: datos fuente del caso."),
        ("A16", "Validacion: controles automáticos de integridad."),
        ("A18", "Preguntas para resolver"),
        ("A19", "1. ¿Cuál es la cartera total y qué porcentaje está vencido?"),
        ("A20", "2. ¿Qué clientes superan su cupo de crédito y cuáles son el top 5 por saldo?"),
        ("A21", "3. ¿Cuánto hay en cada rango de aging y cuánto en riesgo >90 días?"),
        ("A22", "4. ¿Qué pagos no han sido aplicados y a qué facturas deberían aplicarse?"),
        ("A23", "5. ¿Qué diferencias hay entre el extracto bancario y los pagos registrados?"),
        ("A24", "6. ¿Qué notas crédito/débito explican el saldo de una factura específica?"),
        ("A25", "7. ¿Qué cliente priorizarías para cobro y por qué?"),
        ("A27", "Cómo usar"),
        ("A28", "1. Revisá el Dashboard y el Aging."),
        ("A29", "2. Elegí un cliente y reconstruí su saldo desde Facturas, Notas y Aplicacion."),
        ("A30", "3. Conciliá los pagos contra el Extracto y explicá cada diferencia."),
        ("A31", "4. Respondé las preguntas con cifras y una acción concreta de cobro."),
        ("A33", "Nota"),
        (
            "A34",
            "Todos los nombres, NIT, facturas, pagos y extractos son ficticios. El caso no usa información real de empresas.",
        ),
    ]
    for celda, texto in textos:
        ws_ini[celda] = texto
    for fila in [6, 9, 18, 27, 33]:
        ws_ini[f"A{fila}"].font = Font(bold=True, size=12, color=NAVY)
        ws_ini[f"A{fila}"].fill = PatternFill("solid", fgColor=LIGHT)
    for fila in range(7, 35):
        ws_ini[f"A{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_ini.merge_cells("A7:I7")
    ws_ini.merge_cells("A34:I34")
    ws_ini.row_dimensions[7].height = 48
    ws_ini.row_dimensions[34].height = 32
    ajustar_anchos(ws_ini, {"A": 24, "B": 18, "C": 16, "D": 14, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12})

    # ---------------------------------------------------------------
    # Clientes
    # ---------------------------------------------------------------
    ws_cli.append(["ClienteID", "Nombre", "NIT", "Ciudad", "CupoCredito", "PlazoDias"])
    for c in clientes:
        ws_cli.append([c["ClienteID"], c["Nombre"], c["NIT"], c["Ciudad"], int(c["CupoCredito"]), int(c["PlazoDias"])])
    encabezado(ws_cli, 1, 6)
    formato_rango(ws_cli, "E2:E21", MONEY)
    formato_rango(ws_cli, "F2:F21", INT, "center")
    borde_rango(ws_cli, "A1:F21")
    ws_cli.freeze_panes = "A2"
    ws_cli.auto_filter.ref = "A1:F21"
    ajustar_anchos(ws_cli, {"A": 12, "B": 32, "C": 18, "D": 16, "E": 15, "F": 10})

    # ---------------------------------------------------------------
    # Facturas
    # ---------------------------------------------------------------
    ws_fact.append(["FacturaID", "ClienteID", "Cliente", "FechaFactura", "FechaVencimiento", "ValorFactura", "Concepto"])
    for i, f in enumerate(facturas, start=2):
        ws_fact.append([
            f["FacturaID"],
            f["ClienteID"],
            f'=IFERROR(INDEX(Clientes!$B$2:$B$21,MATCH(B{i},Clientes!$A$2:$A$21,0)),"")',
            fecha(f["FechaFactura"]),
            fecha(f["FechaVencimiento"]),
            int(f["ValorFactura"]),
            f["Concepto"],
        ])
    encabezado(ws_fact, 1, 7)
    formato_rango(ws_fact, "D2:E51", DATE, "center")
    formato_rango(ws_fact, "F2:F51", MONEY)
    borde_rango(ws_fact, "A1:G51")
    ws_fact.freeze_panes = "A2"
    ws_fact.auto_filter.ref = "A1:G51"
    ajustar_anchos(ws_fact, {"A": 15, "B": 11, "C": 30, "D": 13, "E": 16, "F": 15, "G": 25})

    # ---------------------------------------------------------------
    # Notas
    # ---------------------------------------------------------------
    ws_notas.append(["Tipo", "NotaID", "FacturaID", "ClienteID", "Fecha", "Valor", "Motivo"])
    for n in nc:
        ws_notas.append(["Crédito", n["NotaCreditoID"], n["FacturaID"], n["ClienteID"], fecha(n["Fecha"]), int(n["Valor"]), n["Motivo"]])
    for n in nd:
        ws_notas.append(["Débito", n["NotaDebitoID"], n["FacturaID"], n["ClienteID"], fecha(n["Fecha"]), int(n["Valor"]), n["Motivo"]])
    encabezado(ws_notas, 1, 7)
    formato_rango(ws_notas, "E2:E15", DATE, "center")
    formato_rango(ws_notas, "F2:F15", MONEY)
    borde_rango(ws_notas, "A1:G15")
    ws_notas.freeze_panes = "A2"
    ws_notas.auto_filter.ref = "A1:G15"
    ajustar_anchos(ws_notas, {"A": 11, "B": 15, "C": 15, "D": 11, "E": 12, "F": 14, "G": 32})
    for fila in range(2, 16):
        if ws_notas[f"A{fila}"].value == "Crédito":
            ws_notas[f"A{fila}"].fill = PatternFill("solid", fgColor=SOFT_GREEN)
        elif ws_notas[f"A{fila}"].value == "Débito":
            ws_notas[f"A{fila}"].fill = PatternFill("solid", fgColor=SOFT_ORANGE)

    # ---------------------------------------------------------------
    # Pagos
    # ---------------------------------------------------------------
    ws_pagos.append([
        "PagoID", "ClienteID", "Cliente", "Fecha", "ValorRecibido", "Aplicado", "NoAplicado",
        "Medio", "Banco", "Referencia", "ValorBanco", "DiferenciaBanco", "EstadoConciliacion", "Observacion",
    ])
    for i, p in enumerate(pagos, start=2):
        ws_pagos.append([
            p["PagoID"],
            p["ClienteID"],
            f'=IFERROR(INDEX(Clientes!$B$2:$B$21,MATCH(B{i},Clientes!$A$2:$A$21,0)),"")',
            fecha(p["Fecha"]),
            int(p["ValorRecibido"]),
            f'=SUMIFS(Aplicacion!$E$2:$E$41,Aplicacion!$A$2:$A$41,A{i})',
            f"=E{i}-F{i}",
            p["Medio"],
            p["Banco"],
            p["Referencia"],
            f'=SUMIFS(Extracto!$D$2:$D$36,Extracto!$B$2:$B$36,J{i})',
            f"=K{i}-E{i}",
            f'=IF(K{i}=0,"Pago sin banco",IF(ABS(L{i})<1,"Conciliado","Diferencia de valor"))',
            p["Observacion"],
        ])
    encabezado(ws_pagos, 1, 14)
    formato_rango(ws_pagos, "D2:D36", DATE, "center")
    formato_rango(ws_pagos, "E2:G36", MONEY)
    formato_rango(ws_pagos, "K2:L36", MONEY)
    borde_rango(ws_pagos, "A1:N36")
    ws_pagos.freeze_panes = "A2"
    ws_pagos.auto_filter.ref = "A1:N36"
    ajustar_anchos(ws_pagos, {"A": 14, "B": 11, "C": 28, "D": 12, "E": 15, "F": 14, "G": 14, "H": 13, "I": 14, "J": 14, "K": 14, "L": 15, "M": 20, "N": 32})
    for rango in ["G2:G36"]:
        for row in ws_pagos[rango]:
            for c in row:
                c.fill = PatternFill("solid", fgColor=SOFT_AMBER if (c.value or 0) not in (0, None) else WHITE)

    # ---------------------------------------------------------------
    # Aplicacion
    # ---------------------------------------------------------------
    ws_apl.append(["PagoID", "FacturaID", "ClienteID", "FechaAplicacion", "ValorAplicado"])
    for a in aplicaciones:
        ws_apl.append([a["PagoID"], a["FacturaID"], a["ClienteID"], fecha(a["FechaAplicacion"]), int(a["ValorAplicado"])])
    encabezado(ws_apl, 1, 5)
    formato_rango(ws_apl, "D2:D41", DATE, "center")
    formato_rango(ws_apl, "E2:E41", MONEY)
    borde_rango(ws_apl, "A1:E41")
    ws_apl.freeze_panes = "A2"
    ws_apl.auto_filter.ref = "A1:E41"
    ajustar_anchos(ws_apl, {"A": 14, "B": 15, "C": 11, "D": 16, "E": 15})

    # ---------------------------------------------------------------
    # Extracto
    # ---------------------------------------------------------------
    ws_ext.append(["Fecha", "Referencia", "Descripcion", "Valor", "Tipo"])
    for e in extracto:
        ws_ext.append([fecha(e["Fecha"]), e["Referencia"], e["Descripcion"], int(e["Valor"]), e["Tipo"]])
    encabezado(ws_ext, 1, 5)
    formato_rango(ws_ext, "A2:A36", DATE, "center")
    formato_rango(ws_ext, "D2:D36", MONEY)
    borde_rango(ws_ext, "A1:E36")
    ws_ext.freeze_panes = "A2"
    ws_ext.auto_filter.ref = "A1:E36"
    ajustar_anchos(ws_ext, {"A": 12, "B": 16, "C": 36, "D": 15, "E": 12})

    # ---------------------------------------------------------------
    # Cartera
    # ---------------------------------------------------------------
    ws_cart.append([
        "FacturaID", "ClienteID", "Cliente", "FechaFactura", "FechaVencimiento", "ValorFactura",
        "NotasCredito", "NotasDebito", "PagosAplicados", "Saldo", "DiasVencidos", "Estado", "Rango",
    ])
    for i, f in enumerate(facturas, start=2):
        ws_cart.append([
            f["FacturaID"],
            f["ClienteID"],
            f'=IFERROR(INDEX(Clientes!$B$2:$B$21,MATCH(B{i},Clientes!$A$2:$A$21,0)),"")',
            fecha(f["FechaFactura"]),
            fecha(f["FechaVencimiento"]),
            int(f["ValorFactura"]),
            f'=SUMIFS(Notas!$F$2:$F$15,Notas!$C$2:$C$15,A{i},Notas!$A$2:$A$15,"Crédito")',
            f'=SUMIFS(Notas!$F$2:$F$15,Notas!$C$2:$C$15,A{i},Notas!$A$2:$A$15,"Débito")',
            f'=SUMIFS(Aplicacion!$E$2:$E$41,Aplicacion!$B$2:$B$41,A{i})',
            f"=F{i}-G{i}+H{i}-I{i}",
            f"=MAX(0,Inicio!$B$4-E{i})",
            f'=IF(J{i}<=0,"Pagada",IF(K{i}=0,"Vigente","Vencida"))',
            f'=IF(J{i}<=0,"Pagada",IF(K{i}=0,"Vigente",IF(K{i}<=30,"0-30",IF(K{i}<=60,"31-60",IF(K{i}<=90,"61-90",">90")))))',
        ])
    encabezado(ws_cart, 1, 13)
    formato_rango(ws_cart, "D2:E51", DATE, "center")
    formato_rango(ws_cart, "F2:J51", MONEY)
    formato_rango(ws_cart, "K2:K51", INT, "center")
    borde_rango(ws_cart, "A1:M51")
    ws_cart.freeze_panes = "C2"
    ws_cart.auto_filter.ref = "A1:M51"
    ajustar_anchos(ws_cart, {"A": 15, "B": 11, "C": 30, "D": 12, "E": 16, "F": 15, "G": 14, "H": 13, "I": 15, "J": 15, "K": 11, "L": 11, "M": 10})

    # Colores por estado de cartera.
    for value, color in [("Pagada", SOFT_GREEN), ("Vigente", SOFT_BLUE), ("0-30", SOFT_AMBER), ("31-60", SOFT_ORANGE), ("61-90", SOFT_RED), (">90", "F5B7B1")]:
        ws_cart.conditional_formatting.add(
            f"M2:M51",
            FormulaRule(formula=[f'$M2="{value}"'], fill=PatternFill("solid", fgColor=color)),
        )
    for value, color in [("Pagada", SOFT_GREEN), ("Vigente", SOFT_BLUE), ("Vencida", SOFT_RED)]:
        ws_cart.conditional_formatting.add(
            f"L2:L51",
            FormulaRule(formula=[f'$L2="{value}"'], fill=PatternFill("solid", fgColor=color)),
        )

    # ---------------------------------------------------------------
    # Aging
    # ---------------------------------------------------------------
    ws_aging.append([
        "ClienteID", "Cliente", "Cupo", "CarteraTotal", "Vigente", "0-30", "31-60", "61-90", ">90",
        "SaldoVencido", "%Cupo", "Semaforo", "Ranking",
    ])
    for i, c in enumerate(clientes, start=2):
        ws_aging.append([
            c["ClienteID"],
            f'=IFERROR(INDEX(Clientes!$B$2:$B$21,MATCH(A{i},Clientes!$A$2:$A$21,0)),"")',
            f'=IFERROR(INDEX(Clientes!$E$2:$E$21,MATCH(A{i},Clientes!$A$2:$A$21,0)),0)',
            f'=SUMIFS(Cartera!$J$2:$J$51,Cartera!$B$2:$B$51,A{i})',
            f'=SUMPRODUCT((Cartera!$B$2:$B$51=A{i})*(Cartera!$L$2:$L$51="Vigente")*Cartera!$J$2:$J$51)',
            f'=SUMPRODUCT((Cartera!$B$2:$B$51=A{i})*(Cartera!$M$2:$M$51="0-30")*Cartera!$J$2:$J$51)',
            f'=SUMPRODUCT((Cartera!$B$2:$B$51=A{i})*(Cartera!$M$2:$M$51="31-60")*Cartera!$J$2:$J$51)',
            f'=SUMPRODUCT((Cartera!$B$2:$B$51=A{i})*(Cartera!$M$2:$M$51="61-90")*Cartera!$J$2:$J$51)',
            f'=SUMPRODUCT((Cartera!$B$2:$B$51=A{i})*(Cartera!$M$2:$M$51=">90")*Cartera!$J$2:$J$51)',
            f"=D{i}-E{i}",
            f"=IF(C{i}=0,0,D{i}/C{i})",
            f'=IF(K{i}>1,"Rojo",IF(K{i}>=0.8,"Amarillo","Verde"))',
            f"=RANK(D{i},$D$2:$D$21,0)+COUNTIF($D$2:D{i},D{i})-1",
        ])
    encabezado(ws_aging, 1, 13)
    formato_rango(ws_aging, "C2:J21", MONEY)
    formato_rango(ws_aging, "K2:K21", PCT, "center")
    formato_rango(ws_aging, "M2:M21", INT, "center")
    borde_rango(ws_aging, "A1:M21")
    ws_aging.freeze_panes = "C2"
    ws_aging.auto_filter.ref = "A1:M21"
    ajustar_anchos(ws_aging, {"A": 11, "B": 30, "C": 14, "D": 15, "E": 14, "F": 13, "G": 13, "H": 13, "I": 13, "J": 15, "K": 10, "L": 11, "M": 9})
    for value, color in [("Rojo", SOFT_RED), ("Amarillo", SOFT_AMBER), ("Verde", SOFT_GREEN)]:
        ws_aging.conditional_formatting.add(
            "L2:L21",
            FormulaRule(formula=[f'$L2="{value}"'], fill=PatternFill("solid", fgColor=color)),
        )

    # Fila de totales del aging.
    ws_aging.cell(row=22, column=1, value="TOTAL")
    for col in range(3, 11):
        letra = get_column_letter(col)
        ws_aging.cell(row=22, column=col, value=f"=SUM({letra}2:{letra}21)").number_format = MONEY
    ws_aging.cell(row=22, column=11, value="=IF(C22=0,0,D22/C22)").number_format = PCT
    for col in range(1, 14):
        c = ws_aging.cell(row=22, column=col)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.font = Font(bold=True, color=NAVY)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center" if col > 2 else "left")

    # ---------------------------------------------------------------
    # Conciliacion
    # ---------------------------------------------------------------
    ws_conc.append(["Referencia", "PagoID", "FechaPago", "ValorPago", "FechaBanco", "ValorBanco", "Diferencia", "Estado", "Observacion"])
    for i, p in enumerate(pagos, start=2):
        ws_conc.append([
            f"=Pagos!J{i}",
            f"=Pagos!A{i}",
            f"=Pagos!D{i}",
            f"=Pagos!E{i}",
            f'=IFERROR(INDEX(Extracto!$A$2:$A$36,MATCH(A{i},Extracto!$B$2:$B$36,0)),"")',
            f'=SUMIFS(Extracto!$D$2:$D$36,Extracto!$B$2:$B$36,A{i})',
            f"=F{i}-D{i}",
            f'=IF(F{i}=0,"Pago sin banco",IF(ABS(G{i})<1,"Conciliado","Diferencia de valor"))',
            f'=IF(H{i}="Conciliado","",IF(H{i}="Pago sin banco","Registrado en cartera, no en banco",IF(H{i}="Diferencia de valor","Diferencia entre recaudo y extracto","")))',
        ])
    # Fila de banco sin pago.
    i = 37
    ws_conc.append([
        "DEP-2026-999",
        "",
        "",
        0,
        f'=IFERROR(INDEX(Extracto!$A$2:$A$36,MATCH(A{i},Extracto!$B$2:$B$36,0)),"")',
        f'=SUMIFS(Extracto!$D$2:$D$36,Extracto!$B$2:$B$36,A{i})',
        f"=F{i}-D{i}",
        "Banco sin pago",
        "Depósito no identificado",
    ])
    ws_conc.append(["TOTAL", "", "", "=SUM(D2:D37)", "", "=SUM(F2:F37)", "=SUM(G2:G37)", "", ""])
    encabezado(ws_conc, 1, 9)
    formato_rango(ws_conc, "C2:C37", DATE, "center")
    formato_rango(ws_conc, "E2:E37", DATE, "center")
    formato_rango(ws_conc, "D2:D38", MONEY)
    formato_rango(ws_conc, "F2:G38", MONEY)
    borde_rango(ws_conc, "A1:I38")
    ws_conc.freeze_panes = "A2"
    ws_conc.auto_filter.ref = "A1:I37"
    ajustar_anchos(ws_conc, {"A": 18, "B": 14, "C": 12, "D": 15, "E": 12, "F": 15, "G": 15, "H": 20, "I": 38})
    for value, color in [("Conciliado", SOFT_GREEN), ("Diferencia de valor", SOFT_AMBER), ("Pago sin banco", SOFT_RED), ("Banco sin pago", SOFT_RED)]:
        ws_conc.conditional_formatting.add(
            "H2:H37",
            FormulaRule(formula=[f'$H2="{value}"'], fill=PatternFill("solid", fgColor=color)),
        )
    for col in range(1, 10):
        c = ws_conc.cell(row=38, column=col)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.font = Font(bold=True, color=NAVY)
        c.border = BORDER

    # ---------------------------------------------------------------
    # Dashboard
    # ---------------------------------------------------------------
    ws_dash.sheet_view.showGridLines = False
    titulo(ws_dash, "Dashboard de cartera — 20 clientes", "A1:I1", 22)
    subtitulo(ws_dash, "Distribuciones Andinas S.A.S. · Datos ficticios · Moneda COP", "A2:I2")
    ws_dash["A3"] = "Corte:"
    ws_dash["A3"].font = Font(bold=True, color=NAVY)
    ws_dash["B3"] = "=Inicio!B4"
    ws_dash["B3"].number_format = DATE
    ws_dash["B3"].font = Font(bold=True, color=TEAL)
    ws_dash["G3"] = "Fuente:"
    ws_dash["G3"].font = Font(bold=True, color=NAVY)
    ws_dash["H3"] = "Facturas, Notas, Pagos y Extracto"
    ws_dash["H3"].font = Font(italic=True, color="5F6F7F")

    # Tarjetas KPI.
    cards = [
        ("A5", "Cartera total", "B5", "=SUM(Cartera!J2:J51)", MONEY),
        ("D5", "Cartera vencida", "E5", '=SUMIFS(Cartera!J2:J51,Cartera!L2:L51,"Vencida")', MONEY),
        ("G5", "% vencida", "H5", "=IF(B5=0,0,E5/B5)", PCT),
        ("A6", "Recaudo", "B6", "=SUM(Pagos!E2:E36)", MONEY),
        ("D6", "Pagos no aplicados", "E6", "=SUM(Pagos!G2:G36)", MONEY),
        ("G6", "Clientes con saldo", "H6", '=COUNTIF(Aging!D2:D21,">0")', INT),
        ("A7", "Facturado", "B7", "=SUM(Facturas!F2:F51)", MONEY),
        ("D7", "Notas crédito", "E7", '=SUMIFS(Notas!F2:F15,Notas!A2:A15,"Crédito")', MONEY),
        ("G7", "Notas débito", "H7", '=SUMIFS(Notas!F2:F15,Notas!A2:A15,"Débito")', MONEY),
        ("A8", "Cupo total", "B8", "=SUM(Clientes!E2:E21)", MONEY),
        ("D8", "% cupo usado", "E8", "=IF(B8=0,0,B5/B8)", PCT),
        ("G8", "Clientes sobre cupo", "H8", '=COUNTIF(Aging!L2:L21,"Rojo")', INT),
    ]
    for label_cell, label, value_cell, formula, fmt in cards:
        ws_dash[label_cell] = label
        ws_dash[label_cell].font = Font(size=9, bold=True, color="5F6F7F")
        ws_dash[label_cell].fill = PatternFill("solid", fgColor=LIGHT)
        ws_dash[label_cell].alignment = Alignment(vertical="center")
        ws_dash[value_cell] = formula
        ws_dash[value_cell].font = Font(size=14, bold=True, color=NAVY)
        ws_dash[value_cell].number_format = fmt
        ws_dash[value_cell].alignment = Alignment(vertical="center")
    for r in range(5, 9):
        ws_dash.row_dimensions[r].height = 28
    for rng in ["A5:B8", "D5:E8", "G5:H8"]:
        for row in ws_dash[rng]:
            for c in row:
                c.border = Border(bottom=Side(style="thin", color=GRAY))

    # Tabla aging.
    ws_dash["A10"] = "Aging de cartera"
    ws_dash["A10"].font = Font(size=12, bold=True, color=NAVY)
    ws_dash["A10"].fill = PatternFill("solid", fgColor=LIGHT)
    ws_dash.merge_cells("A10:C10")
    ws_dash.append([])
    # Fila 11: encabezados manuales (append cambió la última fila; escribimos directo).
    for col, txt in enumerate(["Rango", "Saldo", "% cartera"], start=1):
        c = ws_dash.cell(row=11, column=col, value=txt)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    aging_rows = [
        ("Vigente", '=SUMPRODUCT((Cartera!M2:M51="Vigente")*Cartera!J2:J51)'),
        ("0-30", '=SUMPRODUCT((Cartera!M2:M51="0-30")*Cartera!J2:J51)'),
        ("31-60", '=SUMPRODUCT((Cartera!M2:M51="31-60")*Cartera!J2:J51)'),
        ("61-90", '=SUMPRODUCT((Cartera!M2:M51="61-90")*Cartera!J2:J51)'),
        (">90", '=SUMPRODUCT((Cartera!M2:M51=">90")*Cartera!J2:J51)'),
    ]
    for idx, (label, formula) in enumerate(aging_rows, start=12):
        ws_dash.cell(row=idx, column=1, value=label)
        ws_dash.cell(row=idx, column=2, value=formula).number_format = MONEY
        ws_dash.cell(row=idx, column=3, value=f"=IF($B$17=0,0,B{idx}/$B$17)").number_format = PCT
    ws_dash.cell(row=17, column=1, value="Total")
    ws_dash.cell(row=17, column=2, value="=SUM(B12:B16)").number_format = MONEY
    ws_dash.cell(row=17, column=3, value="=SUM(C12:C16)").number_format = PCT
    for r in range(12, 18):
        for c in range(1, 4):
            cell = ws_dash.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if c > 1 else "left")
    for c in range(1, 4):
        ws_dash.cell(row=17, column=c).fill = PatternFill("solid", fgColor=LIGHT)
        ws_dash.cell(row=17, column=c).font = Font(bold=True, color=NAVY)

    # Top 5 clientes.
    ws_dash["E10"] = "Top 5 clientes por saldo"
    ws_dash["E10"].font = Font(size=12, bold=True, color=NAVY)
    ws_dash["E10"].fill = PatternFill("solid", fgColor=LIGHT)
    ws_dash.merge_cells("E10:I10")
    for col, txt in enumerate(["Cliente", "Saldo", "% cupo", "Semáforo"], start=5):
        c = ws_dash.cell(row=11, column=col, value=txt)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    for k in range(1, 6):
        r = 11 + k
        ws_dash.cell(row=r, column=5, value=f"=INDEX(Aging!$B$2:$B$21,MATCH({k},Aging!$M$2:$M$21,0))")
        ws_dash.cell(row=r, column=6, value=f"=INDEX(Aging!$D$2:$D$21,MATCH({k},Aging!$M$2:$M$21,0))").number_format = MONEY
        ws_dash.cell(row=r, column=7, value=f"=INDEX(Aging!$K$2:$K$21,MATCH({k},Aging!$M$2:$M$21,0))").number_format = PCT
        ws_dash.cell(row=r, column=8, value=f"=INDEX(Aging!$L$2:$L$21,MATCH({k},Aging!$M$2:$M$21,0))")
        for c in range(5, 9):
            ws_dash.cell(row=r, column=c).border = BORDER
        ws_dash.cell(row=r, column=6).alignment = Alignment(horizontal="right")
        ws_dash.cell(row=r, column=7).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=r, column=8).alignment = Alignment(horizontal="center")
    ws_dash.conditional_formatting.add(
        "H12:H16",
        FormulaRule(formula=['$H12="Rojo"'], fill=PatternFill("solid", fgColor=SOFT_RED)),
    )
    ws_dash.conditional_formatting.add(
        "H12:H16",
        FormulaRule(formula=['$H12="Amarillo"'], fill=PatternFill("solid", fgColor=SOFT_AMBER)),
    )
    ws_dash.conditional_formatting.add(
        "H12:H16",
        FormulaRule(formula=['$H12="Verde"'], fill=PatternFill("solid", fgColor=SOFT_GREEN)),
    )

    # Datos para pie.
    ws_dash["A20"] = "Estado"
    ws_dash["B20"] = "Saldo"
    ws_dash["A21"] = "Vigente"
    ws_dash["B21"] = "=B12"
    ws_dash["A22"] = "Vencida"
    ws_dash["B22"] = "=SUM(B13:B16)"
    for r in range(20, 23):
        for c in range(1, 3):
            ws_dash.cell(row=r, column=c).border = BORDER
    ws_dash["B21"].number_format = MONEY
    ws_dash["B22"].number_format = MONEY

    # Resumen de conciliación.
    ws_dash["E20"] = "Conciliación bancaria"
    ws_dash["E20"].font = Font(size=12, bold=True, color=NAVY)
    ws_dash["E20"].fill = PatternFill("solid", fgColor=LIGHT)
    ws_dash.merge_cells("E20:I20")
    conc_rows = [
        ("Pagos registrados", "=SUM(Pagos!E2:E36)", MONEY),
        ("Extracto bancario", "=SUM(Extracto!D2:D36)", MONEY),
        ("Diferencia", "=F22-F21", MONEY),
        ("Pagos sin banco", '=COUNTIF(Conciliacion!H2:H37,"Pago sin banco")', INT),
        ("Banco sin pago", '=COUNTIF(Conciliacion!H2:H37,"Banco sin pago")', INT),
        ("Diferencias de valor", '=COUNTIF(Conciliacion!H2:H37,"Diferencia de valor")', INT),
        ("Pagos conciliados", '=COUNTIF(Conciliacion!H2:H37,"Conciliado")', INT),
    ]
    for idx, (label, formula, fmt) in enumerate(conc_rows, start=21):
        ws_dash.cell(row=idx, column=5, value=label).font = Font(size=9, color="5F6F7F")
        cell = ws_dash.cell(row=idx, column=6, value=formula)
        cell.number_format = fmt
        cell.font = Font(bold=True, color=NAVY)
        for c in range(5, 7):
            ws_dash.cell(row=idx, column=c).border = BORDER

    # Charts.
    chart_aging = BarChart()
    chart_aging.type = "col"
    chart_aging.style = 10
    chart_aging.title = "Cartera por rango de vencimiento"
    chart_aging.y_axis.title = "Saldo COP"
    chart_aging.y_axis.numFmt = '"$"#,##0'
    data = Reference(ws_dash, min_col=2, min_row=11, max_row=16)
    cats = Reference(ws_dash, min_col=1, min_row=12, max_row=16)
    chart_aging.add_data(data, titles_from_data=True)
    chart_aging.set_categories(cats)
    chart_aging.height = 7.5
    chart_aging.width = 11
    chart_aging.legend = None
    chart_aging.dataLabels = DataLabelList()
    chart_aging.dataLabels.showVal = True
    ws_dash.add_chart(chart_aging, "A25")

    chart_top = BarChart()
    chart_top.type = "bar"
    chart_top.style = 10
    chart_top.title = "Top 5 clientes por saldo"
    chart_top.x_axis.title = "Saldo COP"
    chart_top.x_axis.numFmt = '"$"#,##0'
    data_top = Reference(ws_dash, min_col=6, min_row=11, max_row=16)
    cats_top = Reference(ws_dash, min_col=5, min_row=12, max_row=16)
    chart_top.add_data(data_top, titles_from_data=True)
    chart_top.set_categories(cats_top)
    chart_top.height = 7.5
    chart_top.width = 13
    chart_top.legend = None
    chart_top.dataLabels = DataLabelList()
    chart_top.dataLabels.showVal = True
    ws_dash.add_chart(chart_top, "E25")

    chart_pie = PieChart()
    chart_pie.title = "Composición de la cartera"
    data_pie = Reference(ws_dash, min_col=2, min_row=20, max_row=22)
    cats_pie = Reference(ws_dash, min_col=1, min_row=21, max_row=22)
    chart_pie.add_data(data_pie, titles_from_data=True)
    chart_pie.set_categories(cats_pie)
    chart_pie.height = 7.5
    chart_pie.width = 10
    chart_pie.dataLabels = DataLabelList()
    chart_pie.dataLabels.showPercent = True
    ws_dash.add_chart(chart_pie, "A43")

    # Nota al pie.
    ws_dash["E43"] = "Lectura rápida"
    ws_dash["E43"].font = Font(bold=True, color=NAVY)
    ws_dash["E44"] = "La cartera vencida se concentra en los rangos 0-30 y >90. Revisá los clientes en rojo: superan el cupo asignado."
    ws_dash["E44"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_dash.merge_cells("E44:I46")
    ws_dash["E44"].font = Font(size=10, color="5F6F7F")

    ajustar_anchos(ws_dash, {"A": 22, "B": 16, "C": 12, "D": 3, "E": 28, "F": 15, "G": 10, "H": 12, "I": 8})
    ws_dash.freeze_panes = "A5"
    ws_dash.page_setup.orientation = "landscape"
    ws_dash.page_setup.fitToWidth = 1
    ws_dash.page_setup.fitToHeight = 0
    ws_dash.sheet_properties.pageSetUpPr.fitToPage = True
    ws_dash.print_area = "A1:I60"

    # ---------------------------------------------------------------
    # Validacion
    # ---------------------------------------------------------------
    titulo(ws_val, "Validación de integridad", "A1:C1", 18)
    ws_val.append([])
    ws_val.append(["Chequeo", "Resultado", "Detalle"])
    encabezado(ws_val, 3, 3)
    checks = [
        ("Pagos no aplicados no negativos", '=IF(COUNTIF(Pagos!G2:G36,"<0")=0,"OK","Revisar")', "Ningún pago tiene valor aplicado mayor al recibido."),
        ("Cartera = facturas - NC + ND - aplicado", '=IF(ABS(SUM(Cartera!J2:J51)-(SUM(Facturas!F2:F51)-SUMIFS(Notas!F2:F15,Notas!A2:A15,"Crédito")+SUMIFS(Notas!F2:F15,Notas!A2:A15,"Débito")-SUM(Aplicacion!E2:E41)))<1,"OK","Revisar")', "Cuadre del saldo total."),
        ("Aging = cartera total", '=IF(ABS(SUM(Aging!D2:D21)-SUM(Cartera!J2:J51))<1,"OK","Revisar")', "La suma por cliente coincide con el detalle."),
        ("Recaudo = aplicado + no aplicado", '=IF(ABS(SUM(Pagos!E2:E36)-(SUM(Pagos!F2:F36)+SUM(Pagos!G2:G36)))<1,"OK","Revisar")', "Todo el recaudo está clasificado."),
        ("Conciliación = banco - pagos", '=IF(ABS(SUM(Conciliacion!G2:G37)-(SUM(Extracto!D2:D36)-SUM(Pagos!E2:E36)))<1,"OK","Revisar")', "La diferencia está explicada por las partidas."),
        ("Sin saldos de factura negativos", '=IF(COUNTIF(Cartera!J2:J51,"<-1")=0,"OK","Revisar")', "Ninguna factura queda sobrepagada."),
        ("Todos los clientes de facturas existen", '=IF(COUNTIF(Facturas!C2:C51,"")=0,"OK","Revisar")', "No hay clientes huérfanos."),
        ("Ranking de clientes completo", '=IF(COUNTIF(Aging!M2:M21,"")=0,"OK","Revisar")', "Cada cliente tiene posición única."),
    ]
    for idx, (nombre, formula, detalle) in enumerate(checks, start=4):
        ws_val.cell(row=idx, column=1, value=nombre)
        ws_val.cell(row=idx, column=2, value=formula)
        ws_val.cell(row=idx, column=3, value=detalle)
        for c in range(1, 4):
            ws_val.cell(row=idx, column=c).border = BORDER
        ws_val.cell(row=idx, column=2).alignment = Alignment(horizontal="center")
    ws_val.conditional_formatting.add(
        "B4:B11",
        FormulaRule(formula=['$B4="OK"'], fill=PatternFill("solid", fgColor=SOFT_GREEN)),
    )
    ws_val.conditional_formatting.add(
        "B4:B11",
        FormulaRule(formula=['$B4="Revisar"'], fill=PatternFill("solid", fgColor=SOFT_RED)),
    )
    ajustar_anchos(ws_val, {"A": 40, "B": 12, "C": 65})
    ws_val.freeze_panes = "A4"

    # ---------------------------------------------------------------
    # Ajustes finales
    # ---------------------------------------------------------------
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(SALIDA)
    print(f"OK: {SALIDA}")
    print(f"Hojas: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
